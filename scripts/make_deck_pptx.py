#!/usr/bin/env python
"""Native, editable PowerPoint version of the zemi deck (deck_week.tex).

Real text boxes, bullets, native tables, unicode math (editable), the
figures as images, the two audio clips embedded on the concept slide,
and the speaker notes in each slide's notes pane. Mirrors the 22-slide
beamer deck's content; styling is deliberately plain so Ray can restyle.

    python scripts/make_deck_pptx.py   (score-bundle env, repo root)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

INK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x00, 0x72, 0xB2)
MUTED = RGBColor(0x6B, 0x72, 0x80)
FIG = "docs/thesis/figures/"
W, H = Inches(13.33), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H


def slide(title, bullets=(), eq=None, fig=None, fig_w=9.5, note="",
          small=None, audio=()):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if title:
        tb = s.shapes.add_textbox(Inches(0.45), Inches(0.25),
                                  Inches(12.4), Inches(0.75))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, INK
    y = 1.2
    if eq:
        tb = s.shapes.add_textbox(Inches(0.9), Inches(y), Inches(11.5),
                                  Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = eq
        p.font.size, p.font.color.rgb = Pt(20), BLUE
        y += 0.85
    if bullets:
        tb = s.shapes.add_textbox(Inches(0.9), Inches(y), Inches(11.6),
                                  Inches(4.8))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            lvl = 0
            if b.startswith("  "):
                lvl, b = 1, b.strip()
            p.text = ("• " if lvl == 0 else "– ") + b
            p.level = lvl
            p.font.size = Pt(20 if lvl == 0 else 17)
            p.font.color.rgb = INK
            p.space_after = Pt(8)
        y += 0.5 + 0.55 * len(bullets)
    if fig:
        pic = s.shapes.add_picture(FIG + fig, 0, 0, width=Inches(fig_w))
        pic.left = int((W - pic.width) / 2)
        pic.top = Inches(min(y, 7.0 - pic.height / 914400))
    if small:
        tb = s.shapes.add_textbox(Inches(0.9), Inches(6.7), Inches(11.6),
                                  Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = small
        p.font.size, p.font.color.rgb = Pt(13), MUTED
        tb.text_frame.word_wrap = True
    for x, wav in audio:
        s.shapes.add_movie(wav, Inches(x), Inches(6.55), Inches(0.65),
                           Inches(0.65), mime_type="audio/wav")
    if note:
        s.notes_slide.notes_text_frame.text = note
    return s


# 1 title
s = slide(None)
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.8))
for i, (txt, size, bold) in enumerate([
        ("Expressive Performance as a Gaussian Process on a Musical "
         "Score Graph", 32, True),
        ("Raynaldi Lalang — Kyoto University — October 2026", 18, False)]):
    p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
    p.text = txt
    p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, INK

# 2 concept + audio
slide("Expressive performance",
      bullets=["The score fixes pitches and note values. Performers add "
               "the rest: when, how long, how loud, and how the pitch "
               "moves inside the note. These deviations are the "
               "expression."],
      fig="deck_concept.png", fig_w=7.8,
      small="flat (score)  |  expressive (a real performance)  — click "
            "to play",
      audio=[(5.3, "docs/slides/audio/flat.wav"),
             (7.3, "docs/slides/audio/expressive.wav")],
      note="Play flat, then expressive. Same bars, sine synthesis, so "
           "timing and dynamics carry the contrast. Not transcription: "
           "the score is known.")

# 3 thesis
slide("The thesis, and the three problems",
      bullets=["Thesis: one Gaussian process on the score graph turns a "
               "known score and a performance into per-note expressive "
               "variables, each with an error bar you can trust.",
               "Piano: timing, articulation, velocity, read directly "
               "from aligned MIDI. Are the error bars calibrated?",
               "Strings and winds: six variables, now with intonation "
               "and vibrato, estimated from audio, noisily. Does the "
               "same prior still work?",
               "Audio directly: no pitch tracker at all. Can the "
               "waveform itself be the observation?"],
      note="Say the thesis sentence verbatim. One problem per phase; "
           "every later slide answers one of the three.")

# 4 approach
slide("Approach: one prior, three kinds of observation",
      bullets=["One GP prior over per-note expressive fields, on a graph "
               "built from the score. Only the observation changes per "
               "phase: MIDI, f0 targets, or the waveform."],
      fig="deck_programme.png", fig_w=9.6,
      note="The shared prior with three swappable likelihood blocks; "
           "badges show status.")

# 5 graph
slide("The model (1): the score graph",
      eq="Wᵢⱼ = exp( −(bᵢ−bⱼ)²/2ℓ_b² "
         "− (pᵢ−pⱼ)²/2ℓ_p² ),   "
         "L_G = D − W",
      bullets=["Nodes: the written notes (pitch, beat onset, duration).",
               "Edges: close in score time and pitch — neighbours are "
               "played similarly."],
      fig="deck_scoregraph.png", fig_w=8.2,
      note="L_G is the graph Laplacian; it carries the score geometry.")

# 6 coupled GP
slide("The model (2): a GP on the graph, channels coupled",
      eq="K_G(s) = U g(ν; s) Uᵀ,  g(ν; s) = 1/(1+sν)"
         "   ;   Cov[y] = B ⊗ K_G(s)",
      bullets=["Shape the Laplacian spectrum with one parameter s: "
               "smooth-on-the-graph patterns keep their variance, rough "
               "ones are shrunk (graph Matérn).",
               "The channels are correlated: a learned coupling matrix B "
               "shares the graph, so an observed channel at one note "
               "informs the others at neighbouring notes."],
      note="U, nu: eigenvectors/values of the Laplacian. B is learned "
           "per piece.")

# 7 covariance
slide("The model (3): the covariance of one performance",
      eq="K = B⊗K_G(s)  +  Σ_f diag(c_f)⊗X_f X_fᵀ  +  "
         "diag(ς²)⊗I",
      bullets=["Graph term + side-information terms + noise: the whole "
               "model is one per-piece covariance matrix.",
               "Feature weights are re-inferred for every piece — the "
               "measured source of the accuracy gain.",
               "Sixteen numbers per piece, fit by exact evidence; "
               "prediction is conjugate, closed form."],
      note="X_f: per-note features (score descriptors, music-model "
           "embeddings). Baselines are exact special cases.")

# 8 architecture
slide("The model, in one picture", fig="arch_phase1_deck.png", fig_w=10.6,
      small="Blue: score graph → L_G → K_G. Amber: features as "
            "linear kernels. Vermilion: the likelihood, the only block "
            "that changes across the phases.",
      note="Walk the colors left to right.")

# 9 P1 setting
slide("Phase 1: the setting (piano, aligned MIDI)",
      eq="yᵢ = [ τᵢ (onset timing),  log rᵢ "
         "(articulation),  vᵢ (velocity) ]",
      bullets=["A Disklavier records exactly what the pianist "
               "controlled: three variables per note, read directly, no "
               "estimation.",
               "The question: on held-out notes, how good are the "
               "estimates, and can the error bars be trusted?"],
      note="The concept picture was exactly this data.")

# 10 P1 results
slide("Phase 1: results (piano, ASAP)",
      bullets=["Preregistered, scored once on 20 untouched pieces, "
               "against the two-stage pipeline the model contains.",
               "Accuracy: RMSE 0.376 vs 0.393 (pooled over the three "
               "channels), paired per-piece difference significant.",
               "Calibration: the graph's log-score contribution "
               "significant; coverage 0.925 at nominal 0.90.",
               "Attribution: the features carry the mean, the graph "
               "carries the calibration."],
      note="One-shot discipline: the test set was never touched during "
           "development.")

# 11 cents curve
slide("Phase 2: from audio to the cents curve",
      eq="xⱼ = 1200 log₂(f₀ⱼ/440) − "
         "100(pᵢ−69)   (deviation from equal temperament)",
      bullets=["A pitch tracker gives frames: time, f₀, voicing "
               "confidence.",
               "Frame rule: voiced, confidence above the track's lowest "
               "quintile, at least 4 frames per note.",
               "Two curves per note — tracked: pYIN, run by us; ground "
               "truth: the corpus's own annotation of the same stem. "
               "Cleaner, but still a measurement."],
      note="GT = corrected pYIN on isolated stems (46 ms / 10 ms), gross "
           "errors hand-fixed; hence 'quasi-truth'.")

# 12 channels
slide("Phase 2: the pitch-curve channels",
      bullets=["Within one note: a sustained level plus one oscillation. "
               "Three numbers describe it.",
               "Intonation c (cents): the sustained offset from the "
               "written pitch. c > 0 sharp; 100 = a semitone.",
               "Vibrato extent γ (cents): the oscillation's "
               "amplitude.",
               "Vibrato rate f (Hz): its frequency; vibrato lives at "
               "roughly 4–8 Hz.",
               "The bundle carries log γ, log f. No oscillation "
               "found: those cells are missing, c remains."],
      note="Logs: positive quantities, multiplicative errors.")

# 13 data figure
slide("Phase 2: what the data looks like",
      fig="phase2_frames_intro.png", fig_w=11.6,
      small="A: the tracker quantizes and drops frames; the ground "
            "truth shows the oscillation underneath. B: identifiable. "
            "C: refused, cells missing.",
      note="Real development notes, no fits drawn.")

# 14 GP estimator
slide("Estimating the channels with a Gaussian process",
      eq="k(Δt) = w₁ e^(−2π²v₁Δt²) "
         "cos(2πμ₁Δt) + w₂ "
         "e^(−2π²v₂Δt²)",
      bullets=["The cents curve as a Gaussian process, two components: "
               "a damped cosine at the vibrato rate, plus a slow drift "
               "(a special case of the spectral mixture kernel, Wilson "
               "and Adams 2013).",
               "The prior states only where the curve's energy lives: a "
               "band at the rate, plus a drift band at zero.",
               "Fit per note by exact marginal likelihood; variances "
               "from the evidence curvature.",
               "Channels from the posterior: c the realized average, "
               "γ from the vibrato component, f = μ₁."],
      note="Naming, say once: the kernel is a special case of the SM "
           "kernel; the method is ours (structure imposed). We also ran "
           "the SM method as intended on 9,804 notes: exact tie at "
           "curve level, vibrato band found unaided on 18%. The "
           "structure must be imposed and costs nothing.")

# 15 one-note figure
slide("One real note: the curve and its parts",
      fig="sm_estimator_explainer.png", fig_w=11.6,
      small="A: the GP tracks the curve (rigid sine as reference). B: "
            "the posterior split into centre, drift, vibrato. C: the "
            "read-outs.",
      note="The split IS the channel read-out.")

# 16 P2 results
slide("Phase 2: results (strings and winds, URMP)",
      eq="yᵢ = [ c, log γ, log f, ℓ, τ, dᵢᵛⁱᵇ ]"
         "   — six channels, estimator variances used as given",
      bullets=["Preregistered, scored once on 13 held-out pieces. "
               "Δ = full model − no-graph ablation, paired per "
               "piece.",
               "Intonation recovery: ΔRMSE −0.877 cents, "
               "significant.",
               "Vibrato calibration: ΔNLL −2.990 (extent), "
               "−0.564 (rate), both significant.",
               "Coverage 0.88–0.91 at nominal 0.90, all six channels.",
               "Timing calibration: no improvement (−0.030, interval "
               "includes zero). The loose end, taken up after Phase 3."],
      note="Same prior as piano, unchanged.")

# 17 P3
slide("Phase 3: the waveform as the observation",
      eq="x = Φ(z) a + ε   ⇒   p(x|z) = N(x; 0, "
         "ΦΣ_aΦᵀ + σ²I)",
      bullets=["The likelihood block becomes the audio itself; the "
               "harmonic amplitudes are marginalized exactly "
               "(O(mp²), whole tracks exact).",
               "Development studies: intonation at 2.29 cents median "
               "with no pitch tracker; beats the estimator chain on "
               "winds."],
      note="x: audio samples; z: pitch deviations; a: amplitudes, never "
           "estimated.")

# 18 prototype
slide("Running now: the curve straight from the waveform",
      fig="phase3_curve_proto.png", fig_w=11.6,
      small="Prototype (this week): the within-note GP prior meets the "
            "waveform likelihood. No tracker, no estimator. Better than "
            "the pYIN frames on two of three notes.",
      note="Do not cut. Violin 2.5 vs pYIN 3.4 cents, clarinet 1.0 vs "
           "4.6, cello 4.8 vs 2.7 (the miss, say it plainly). Knots "
           "must beat Nyquist for 5-6 Hz vibrato: 16 knots/s.")

# 19 timing
slide("Timing noise is correlated across notes",
      eq="Σᵗᵢⱼ = ς_{τ,i} ς_{τ,j} "
         "ρ^|i−j|   — one parameter ρ, chosen per piece "
         "by the evidence",
      bullets=["Why timing did not calibrate: alignment error is "
               "correlated along score time; a diagonal noise row "
               "cannot represent it.",
               "The evidence detects the correlation on 97–100% of "
               "pieces (median ρ 0.45–0.60).",
               "Calibration improves (ΔNLL −0.111, significant); "
               "timing error falls ~20% (99 → 80 ms).",
               "Reproduced three ways: original masks, fresh masks, "
               "joint refit."],
      note="Say 'correlated timing noise across notes', never bare "
           "'AR(1)'.")

# 20 mechanism
slide("Held-out timing, with and without the correlation",
      fig="corrnoise_tau_dev.png", fig_w=11.6,
      small="Left: paired deltas, three runs. Middle: the chosen "
            "correlation per piece. Right: one track's held-out timing.",
      note="A neighbour's warp error predicts yours; the correlated row "
           "lets the model subtract it.")

# 21 timeline (native table)
s = slide("From here (proposed sequencing)",
          note="Sequencing, not dates; gated by the open questions.")
rows = [("now", "integrate the full-corpus learned-spectrum run (done "
                "this week)"),
        ("next 2 weeks", "freeze the timing-noise registration once the "
                         "pool is decided; batch-evaluate the "
                         "curve-from-waveform prototype"),
        ("November", "decide the channel design (scalars vs curve "
                     "level); run the timing confirmation, one shot"),
        ("December on", "Phase 3 at scale; consolidate thesis chapters")]
tbl = s.shapes.add_table(4, 2, Inches(1.2), Inches(1.5), Inches(10.9),
                         Inches(4.2)).table
tbl.columns[0].width, tbl.columns[1].width = Inches(2.6), Inches(8.3)
for r, (a, b) in enumerate(rows):
    for c, txt in enumerate((a, b)):
        cell = tbl.cell(r, c)
        cell.text = txt
        cell.text_frame.paragraphs[0].font.size = Pt(18)

# 22 open questions
slide("Open questions",
      bullets=["Which corpus to confirm the timing result on: a new one "
               "(Bach10), a disclosed re-use, or both?",
               "Keep the channels as three scalars per note, or move "
               "them to curve level (the prototype's direction)?",
               "Should the new robustness options become defaults?"],
      note="End here; invite discussion.")

prs.save("docs/slides/deck_week.pptx")
print("native pptx:", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
